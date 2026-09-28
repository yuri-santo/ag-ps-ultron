"""Ata formal de reunião em Word (.docx) editável, com o mesmo conteúdo e identidade do PDF.

Usa build_minutes() do render_report, então as duas versões nunca divergem.
Uso: render_docx.py report-input.json saida.docx
"""
import argparse
import io
import json
from pathlib import Path

try:
    from .render_report import ACCENT, ACCENT_DEEP, ACCENT_SOFT, INK, MUTED, NAO_APURADO, _stamp, build_minutes, quem_falou
except ImportError:
    import sys
    sys.path.insert(0, str(Path(__file__).resolve().parent))
    from render_report import ACCENT, ACCENT_DEEP, ACCENT_SOFT, INK, MUTED, NAO_APURADO, _stamp, build_minutes, quem_falou


def _rgb(hex_color):
    from docx.shared import RGBColor
    return RGBColor.from_string(hex_color.lstrip('#').upper())


def logo_png(size=240):
    """Hexágono do Ultron (mesmo ícone do botão do Stream Deck) desenhado com Pillow."""
    from PIL import Image, ImageDraw
    img = Image.new('RGBA', (size, size), (0, 0, 0, 0))
    d = ImageDraw.Draw(img)
    s = size / 24.0

    def pt(x, y):
        return (x * s, y * s)
    hexagon = [pt(12, 2), pt(21, 7), pt(21, 17), pt(12, 22), pt(3, 17), pt(3, 7)]
    d.polygon(hexagon, fill=ACCENT_SOFT, outline=ACCENT_DEEP)
    d.line(hexagon + [hexagon[0]], fill=ACCENT_DEEP, width=max(3, int(1.9 * s)), joint='curve')
    w = max(3, int(1.6 * s))
    d.line([pt(7, 9), pt(12, 12), pt(17, 9)], fill=ACCENT, width=w, joint='curve')
    d.line([pt(12, 12), pt(12, 18)], fill=ACCENT, width=w)
    buf = io.BytesIO()
    img.save(buf, 'PNG')
    buf.seek(0)
    return buf


def render_docx(data, destination):
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt

    ata = build_minutes(data)
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.left_margin = sec.right_margin = Cm(2)
    sec.top_margin, sec.bottom_margin = Cm(1.6), Cm(1.8)
    base = doc.styles['Normal']
    base.font.name = 'Arial'
    base.font.size = Pt(10)
    base.font.color.rgb = _rgb(INK)

    def shade(cell, color):
        tc = cell._tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'), color.lstrip('#'))
        tc.append(shd)

    def heading(text):
        par = doc.add_paragraph()
        par.paragraph_format.space_before = Pt(12)
        par.paragraph_format.space_after = Pt(4)
        run = par.add_run(text)
        run.bold = True
        run.font.size = Pt(13)
        run.font.color.rgb = _rgb(ACCENT_DEEP)
        return par

    def para(text, muted=False, size=10):
        par = doc.add_paragraph()
        par.paragraph_format.space_after = Pt(3)
        run = par.add_run(text)
        run.font.size = Pt(size)
        if muted:
            run.font.color.rgb = _rgb(MUTED)
        return par

    # Cabeçalho: logotipo + ULTRON à esquerda, título à direita (como no modelo)
    head = doc.add_table(rows=1, cols=2)
    head.alignment = WD_TABLE_ALIGNMENT.CENTER
    left, right = head.rows[0].cells
    lp = left.paragraphs[0]
    lp.add_run().add_picture(logo_png(), width=Cm(1.6))
    r = lp.add_run('  ULTRON')
    r.bold = True
    r.font.size = Pt(18)
    r2 = left.add_paragraph().add_run('COMMAND DECK · SAITAMA')
    r2.font.size = Pt(7.5)
    r2.font.color.rgb = _rgb(MUTED)
    rp = right.paragraphs[0]
    rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    t = rp.add_run('ATA FORMAL DE REUNIÃO')
    t.bold = True
    t.font.size = Pt(16)
    t.font.color.rgb = _rgb(ACCENT_DEEP)
    sp = right.add_paragraph()
    sp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    sr = sp.add_run(f"{ata['titulo'][:70]} · {ata['data']}")
    sr.font.size = Pt(8.5)
    sr.font.color.rgb = _rgb(MUTED)

    footer = sec.footer.paragraphs[0]
    fr = footer.add_run('Ultron · ata gerada a partir do registro da reunião; conferir antes de assinar')
    fr.font.size = Pt(7.5)
    fr.font.color.rgb = _rgb(MUTED)

    def lista(nomes):
        return ', '.join(nomes) if nomes else NAO_APURADO

    nome = '' if ata['titulo'].strip().lower() in ('reunião', 'reuniao', '') else f" {ata['titulo']}"
    escopo = f" do {ata['escopo']}" if ata['escopo'] else ''
    heading('Chamada')
    para(f"A reunião{nome}{escopo} foi realizada em {ata['data']}, em {ata['local']}. Começou às {ata['inicio']} "
         f"e foi presidida por {ata['presidente']}, com {ata['secretaria']} no cargo de secretaria.")
    heading('Participantes')
    para(f"Presença registrada no áudio: {lista(ata['presentes'])}.")
    para(f"Convocados pelo convite: {lista(ata['convocados'])}.")
    para(f"Membros ausentes: {NAO_APURADO} (a captura não confirma presença nominal).")
    heading('Aprovação das Atas Anteriores')
    para('Não houve pedido de aprovação de atas anteriores registrado nesta reunião.')
    heading('Relatórios Apresentados')
    linhas = [x for x in ata['relatorios'].splitlines() if x.strip()]
    if linhas:
        para('Chamados e documentos consultados como contexto (não são falas da reunião):')
        for linha in linhas:
            doc.add_paragraph(linha, style='List Bullet')
    else:
        para('Nenhum relatório formal foi registrado nesta reunião.')
    heading('Assuntos Pendentes')
    if ata['pendentes'] or ata['revogadas']:
        for item in ata['pendentes']:
            claim = item.get('claim') if isinstance(item, dict) else str(item)
            question = item.get('question') if isinstance(item, dict) else ''
            doc.add_paragraph(claim + (f' — pergunta sugerida: {question}' if question else ''), style='List Bullet')
        for record in ata['revogadas']:
            doc.add_paragraph(f"Revogado/corrigido: {record.get('text', '')} [{_stamp(record)}]", style='List Bullet')
        if ata['pendentes']:
            para('Itens acima são sugestões do copiloto pendentes de confirmação; não são decisões da reunião.', muted=True, size=8)
    else:
        para('Nenhum assunto pendente registrado.')
    heading('Deliberações')
    if ata['decisoes']:
        for record in ata['decisoes']:
            autor = quem_falou(record, ata['rotulos'])
            doc.add_paragraph(f"{record.get('text', '')} [{_stamp(record)}" + (f'; {autor}' if autor else '') + ']',
                              style='List Bullet')
    else:
        para('Nenhuma decisão confirmada com citação válida.')
    if ata['acoes']:
        sub = para('Encaminhamentos')
        sub.runs[0].bold = True
        sub.runs[0].font.color.rgb = _rgb(ACCENT_DEEP)
        table = doc.add_table(rows=1, cols=4)
        table.style = 'Table Grid'
        for cell, title in zip(table.rows[0].cells, ('Ação', 'Responsável', 'Prazo', 'Registro')):
            cell.text = title
            cell.paragraphs[0].runs[0].bold = True
            cell.paragraphs[0].runs[0].font.color.rgb = _rgb(ACCENT_DEEP)
            shade(cell, ACCENT_SOFT)
        for record in ata['acoes']:
            row = table.add_row().cells
            row[0].text = record.get('text', '')
            row[1].text = record.get('owner') or 'não informado'
            row[2].text = record.get('deadline') or 'não informado'
            row[3].text = f"{_stamp(record)} · {', '.join(record.get('transcript_ids', []))}"
    heading('Encerramento')
    fim = f"às {ata['fim']}." if ata['fim'] != NAO_APURADO else '(horário de encerramento não apurado pelo registro automático).'
    para(f"Nada mais havendo a tratar, a reunião foi encerrada {fim} A presente ata foi lavrada pela secretaria a partir "
         'da gravação e da transcrição da reunião e segue para aprovação dos participantes.')
    para(f"Ata enviada por: Ultron · Ata aprovada por: ______________________ · Sessão {ata['sessao']}", muted=True, size=8)
    sign = doc.add_table(rows=2, cols=2)
    sign.rows[0].cells[0].text = '_' * 34
    sign.rows[0].cells[1].text = '_' * 34
    sign.rows[1].cells[0].text = f"{ata['presidente']}\nPresidente"
    sign.rows[1].cells[1].text = 'Ultron\nSecretaria'

    doc.add_page_break()
    heading('Anexo I — Análise executiva da reunião')
    for line in (ata['report_text'] or 'Análise indisponível.').split('\n'):
        clean = line.strip()
        if clean.startswith('#'):
            sub = para(clean.lstrip('#').strip())
            sub.runs[0].bold = True
            sub.runs[0].font.color.rgb = _rgb(ACCENT_DEEP)
        elif clean.startswith(('- ', '* ')):
            doc.add_paragraph(clean[2:].replace('**', ''), style='List Bullet')
        else:
            para(line.replace('**', ''))
    doc.add_page_break()
    heading('Anexo II — Transcrição integral (horários e origem do áudio)')
    for line in (ata['transcript'] or '[Nenhuma fala transcrita]').split('\n'):
        para(line, size=8.5)
    doc.core_properties.title = 'Ata formal de reunião — ' + ata['titulo']
    doc.core_properties.author = 'Ultron'
    doc.save(str(destination))


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('input')
    parser.add_argument('output')
    args = parser.parse_args()
    render_docx(json.loads(Path(args.input).read_text(encoding='utf-8')), args.output)
